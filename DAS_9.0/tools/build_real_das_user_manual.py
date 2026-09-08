from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"D:\DAS_9.0(1)\DAS_9.0\docs\Real_Das_用户使用说明书.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width.twips))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = width


def set_fixed_table_layout(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(w.twips for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width.twips))
        grid.append(col)
    for i, width in enumerate(widths):
        table.columns[i].width = width


def default_widths(col_count):
    if col_count == 2:
        return [Cm(4.0), Cm(12.2)]
    if col_count == 3:
        return [Cm(2.6), Cm(3.9), Cm(9.7)]
    if col_count == 4:
        return [Cm(3.2), Cm(3.1), Cm(3.2), Cm(6.7)]
    if col_count == 5:
        return [Cm(2.8), Cm(2.6), Cm(2.8), Cm(2.8), Cm(5.2)]
    return [Cm(16.2 / col_count)] * col_count


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True):
    p_pr = paragraph._p.get_or_add_pPr()
    elem = p_pr.find(qn("w:keepNext"))
    if keep and elem is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not keep and elem is not None:
        p_pr.remove(elem)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color="4F81BD"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(doc, text="", style=None, color=None, bold=False, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.clear()
    r = p.add_run(text)
    colors = {1: "16314D", 2: "1F4E79", 3: "35506C"}
    sizes = {1: 17, 2: 14, 3: 12}
    set_font(r, size=sizes.get(level, 12), bold=True, color=colors.get(level, "35506C"))
    set_keep_with_next(p, True)
    return p


def add_note(doc, title, body, fill="EAF3FF", border="9BC2E6"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    set_paragraph_shading(p, fill)
    set_paragraph_left_border(p, border)
    r = p.add_run(title + "\n")
    set_font(r, size=10.5, bold=True, color="16314D")
    r2 = p.add_run(body)
    set_font(r2, size=10.5, color="263747")
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None, header_fill="1F4E79"):
    for idx, row in enumerate(rows, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(4)
        if len(row) == 1:
            r = p.add_run(str(row[0]))
            set_font(r, size=10.2, color="263747")
            continue
        lead = str(row[0])
        r = p.add_run(f"{lead}")
        set_font(r, size=10.2, bold=True, color="1F4E79")
        remaining = []
        for h, v in zip(headers[1:], row[1:]):
            remaining.append(f"{h}：{v}")
        r2 = p.add_run(" — " + "；".join(remaining))
        set_font(r2, size=10.0, color="263747")
    doc.add_paragraph()
    return None


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.9)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = styles[style_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("Real_Das 分布式光纤振动监测软件使用说明")
    set_font(r, size=8.5, color="6B7D90")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("内部培训与现场操作资料")
    set_font(r, size=8.5, color="6B7D90")


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Real_Das 软件使用说明书")
    set_font(r, size=25, bold=True, color="16314D")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("分布式光纤振动监测系统 · 现场用户版")
    set_font(r, size=14, bold=True, color="4F81BD")

    doc.add_paragraph()
    add_note(
        doc,
        "使用前先读这一句",
        "本软件的操作顺序是：先选光纤长度，再选抽取系数，再设监测起止通道，最后设差分距离（标距）。差分距离对应空间分辨率，现场不得设置到低于 3.2 m 的物理标距。",
        fill="EAF3FF",
        border="4F81BD",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("整理依据：当前工程代码与主界面控件\n适用对象：现场操作人员、调试人员、售后培训人员")
    set_font(r, size=10.5, color="44546A")
    doc.add_section(WD_SECTION_START.NEW_PAGE)


def add_quick_start(doc):
    add_heading(doc, "一、推荐操作顺序", 1)
    steps = [
        ("1", "选择光纤长度", "在“光纤长度”下拉框选择 5km、10km、30km 或 50km，然后点击旁边“确定”。软件会同步设置对应脉冲频率。"),
        ("2", "选择抽取系数", "在“抽取系数”中选择抽取倍数后点击“确定”。抽取是为了降低上位机实时处理和显示的数据量。"),
        ("3", "设置监测起始/结束通道", "输入监测起始通道和监测结束通道，分别点击对应“确定”。监测范围越大，数据量越大。"),
        ("4", "设置差分距离（标距）", "输入差分点数后点击“确定”。它决定物理标距/空间分辨率，现场设置不得小于 3.2 m。"),
        ("5", "确认滤波和显示模式", "一般保持“是否滤波”勾选，高通截止默认 1.00 Hz；根据需要切换 RMS 瀑布/单帧、FFT 瀑布/频谱。"),
        ("6", "观察处理耗时", "原始数据、数据重构、解缠绕、数据滤波、RMS计算、fft计算、存储耗时不应变红。红色表示处理跟不上数据进入速度。"),
        ("7", "按需保存/查看数据", "保存单点用于长期保存某个通道波形；保存全部用于保存完整空间范围数据；查看单点/查看保存全部用于离线回看。"),
    ]
    add_table(doc, ["步骤", "操作", "说明"], steps, widths=[Cm(1.1), Cm(3.6), Cm(11.5)])

    add_heading(doc, "二、软件内部数据流程", 1)
    add_paragraph(
        doc,
        "从代码连接关系看，软件的数据链路是：PCIE 原始数据读取 → 数据重构（按差分距离做空间差分）→ 解缠绕 → 保存快照 → 滤波 → RMS 计算与 FFT 计算 → 主界面显示。界面右侧的处理耗时标签正是对这些环节逐段计时。",
    )
    flow_rows = [
        ("原始数据", "receive_data", "从 FPGA/PCIE 读取一帧原始数据。"),
        ("数据重构", "rebuild_data", "按起止通道、抽取系数和差分距离重构出相位矩阵。"),
        ("解缠绕", "unwrap", "对相位数据做连续性修正，并把快照送入保存模块。"),
        ("数据滤波", "Filter", "根据“是否滤波”和高通截止频率处理数据，并抽取音频通道。"),
        ("RMS计算", "rms_calculate", "计算振动定位瀑布图或单帧曲线所需的 RMS 值。"),
        ("fft计算", "fft_calculate", "对当前通道波形做 FFT，更新频谱瀑布图或频谱图。"),
        ("保存数据", "sava_data", "保存单点或保存全部数据，并维护磁盘空间与保存状态。"),
    ]
    add_table(doc, ["界面名称", "代码模块", "作用"], flow_rows, widths=[Cm(2.6), Cm(3.1), Cm(10.5)])


def add_parameter_sections(doc):
    add_heading(doc, "三、光纤长度：先确定量程和基础点距", 1)
    add_paragraph(
        doc,
        "“光纤长度”不是单纯的文字显示，它在代码中对应脉冲频率。点击光纤长度旁边的“确定”后，软件把 pulse_frequency 写入参数管理器，并刷新监测通道范围、RMS 横轴和 FFT 图。",
    )
    rows = [
        ("5km(20Khz)", "20000 Hz", "0.4 m/原始点", "约 11264 行", "数据进入最快，实时压力最高。"),
        ("10km(10Khz)", "10000 Hz", "0.4 m/原始点", "约 23552 行", "默认档位，常用基准。"),
        ("30km(3.333Khz)", "3333 Hz", "1.2 m/原始点", "约 24576 行", "点距变为 10km 的 3 倍。"),
        ("50km(2Khz)", "2000 Hz", "2.0 m/原始点", "约 24576 行", "点距更大，差分点数要人工调小。"),
    ]
    add_table(doc, ["界面选项", "内部频率", "FPGA上传点距", "原始行数", "操作提示"], rows)
    add_note(
        doc,
        "现场规则",
        "先选光纤长度，再设置抽取系数和差分距离。因为不同长度下 FPGA 上传的单点距离不同，后面的通道间距和空间分辨率都会随之变化。",
        fill="FFF7E6",
        border="F4B183",
    )

    add_heading(doc, "四、抽取系数：为什么必须抽取", 1)
    add_paragraph(
        doc,
        "分布式光纤数据量非常大，一帧数据同时包含空间方向大量通道和时间方向大量采样点。如果不抽取，PCIE读取、重构、解缠绕、滤波、RMS、FFT、绘图和保存都会承受很大压力。抽取系数就是在空间方向按间隔取点，减少参与显示与后续计算的通道数量。",
    )
    rows = [
        ("不抽取数据", "1", "保留全部点", "仅用于短距离、低压力或调试。"),
        ("抽取2倍", "2", "每 2 个点取 1 个", "空间显示更稀疏，压力约减半。"),
        ("抽取4倍", "4", "每 4 个点取 1 个", "中等压力配置。"),
        ("抽取8倍", "8", "每 8 个点取 1 个", "默认配置，10km 时约 3.2m/通道。"),
        ("抽取16倍", "16", "每 16 个点取 1 个", "数据压力大或距离范围大时使用。"),
    ]
    add_table(doc, ["界面选项", "内部值", "含义", "建议"], rows)
    add_note(
        doc,
        "重要区别",
        "抽取系数影响“监测通道”和图上位置间距；差分距离影响“标距/空间分辨率”。两者不是一个参数。改变抽取系数不会自动改变差分距离，现场需要人工检查。",
        fill="E2F0D9",
        border="70AD47",
    )

    add_heading(doc, "五、监测起始通道和结束通道", 1)
    add_paragraph(
        doc,
        "监测起始通道和监测结束通道控制软件真正处理和显示的空间范围。输入通道号后点击右侧“确定”，软件会把数值限制在当前光纤长度和抽取系数允许的最大范围内。",
    )
    rows = [
        ("监测起始通道", "lineEdit / start_save", "输入开始处理的位置；小于 0 会被限制为 0。"),
        ("监测结束通道", "lineEdit_2 / end_save", "输入结束处理的位置；必须大于等于起始通道。"),
        ("旁边的“m/通道”", "meterPerRawPoint × extract", "显示当前每个监测通道对应的物理距离。"),
    ]
    add_table(doc, ["界面项", "内部参数", "说明"], rows)
    add_paragraph(
        doc,
        "示例：10km 档位的原始点距为 0.4m，抽取 8 倍时就是 0.4 × 8 = 3.2m/通道。若结束通道显示为 2944，则监测范围大约为 2944 × 3.2m。",
    )

    add_heading(doc, "六、差分距离（标距）：空间分辨率设置", 1)
    add_paragraph(
        doc,
        "差分距离是重构模块做空间差分时使用的点数。代码中它直接作为 FPGA 上传后的原始点偏移，不跟随抽取系数自动变化。物理标距计算公式为：差分距离 × 当前量程下的 FPGA 上传点距。",
    )
    rows = [
        ("5km / 20k", "0.4 m/点", "8", "3.2 m", "不能小于 8。"),
        ("10km / 10k", "0.4 m/点", "8", "3.2 m", "默认推荐。"),
        ("30km / 3.333k", "1.2 m/点", "3", "3.6 m", "2 点只有 2.4m，低于 3.2m，不建议。"),
        ("50km / 2k", "2.0 m/点", "2", "4.0 m", "1 点只有 2m，低于 3.2m，不建议。"),
    ]
    add_table(doc, ["光纤长度", "上传点距", "最小建议差分点数", "对应标距", "说明"], rows)
    add_note(
        doc,
        "空间分辨率底线",
        "从光域能力看，本系统最优空间分辨率按 3.2m 处理。用户可以把差分距离设得更大以获得更稳定、更平滑的结果，但不能把物理标距设置小于 3.2m。小于 3.2m 时界面可能仍能计算，但不代表真实空间分辨率提高。",
        fill="FCE4D6",
        border="C65911",
    )


def add_controls(doc):
    add_heading(doc, "七、主界面按钮和控件说明", 1)
    rows = [
        ("查看单点", "打开单点保存数据回看窗口。选择单点保存文件夹后，显示该通道波形和 FFT 频谱。"),
        ("查看保存全部", "打开保存全部数据回看窗口。先选保存全部文件夹，再选择其中一个 bin 文件，显示完整瀑布图、所选通道波形和所选通道 FFT。"),
        ("通道面板 1/2 的“通道”输入框", "输入要实时查看的通道号。界面会同步显示该通道对应的米数和当前 m/通道。"),
        ("通道面板“确定”", "确认通道号，刷新该通道实时波形。通道 1 同时会更新 monitorPosition。"),
        ("通道面板“保存单点”", "保存当前面板通道的一维时间波形。点击后选择保存目录，软件收到下一帧后自动建立带参数的文件夹。"),
        ("通道面板“结束保存”", "结束当前单点保存。保存期间通道输入会被锁定，避免保存中途改通道导致数据混乱。"),
        ("通道面板“初始化”", "重置当前实时波形图的显示范围。曲线有数据时会自动跟随最近数据，否则回到 0-5 秒、-1 到 1 rad。"),
        ("光纤长度旁“确定”", "把下拉框中的长度写入系统参数，并更新频率、监测范围、RMS/FFT坐标。"),
        ("抽取系数旁“确定”", "把抽取倍数写入系统参数，并刷新通道间距、监测范围和 RMS 显示。"),
        ("监测起始通道“确定”", "确认起始通道；若起始大于结束，软件会把结束通道同步提高到起始通道。"),
        ("监测结束通道“确定”", "确认结束通道；软件会限制它不小于起始通道、不超过当前最大通道。"),
        ("差分距离(标距)“确定”", "确认差分点数，并弹窗显示当前物理标距。输入必须为大于等于 1 的整数；现场还要人工保证物理标距不小于 3.2m。"),
        ("是否滤波", "勾选时启用滤波链路；一般现场保持勾选。取消勾选会切换到应变相关显示，建议只在调试或明确需求下使用。"),
        ("差分相位", "高级选项。开启后对同一空间行做时间方向差分相位处理，用于观察相邻时间采样变化；一般保持关闭。"),
        ("高通截止(Hz) + 滤波确定", "设置高通滤波器截止频率。默认 1.00Hz；必须大于 0 且小于当前频率的一半。"),
        ("保存全部", "保存当前监测范围内的完整二维数据。点击后选择目录，软件会自动建立 data_all_ 开头的参数文件夹。"),
        ("结束保存全部", "停止保存全部数据，并关闭当前保存文件。"),
        ("保存配置", "把当前参数保存为 txt 配置文件，包括光纤长度、抽取、起止通道、差分距离、滤波、高通、线程数等。"),
        ("读取配置", "从 txt 配置文件恢复参数。读取后会刷新界面和图形坐标。"),
        ("开始", "当前代码中主要用于状态栏显示“开始”；真实数据是否进入取决于硬件和 PCIE 数据状态。"),
        ("停止", "当前代码中主要用于状态栏显示“停止”；不作为硬件断电或强制停止保存使用。需要停止保存时请使用对应结束保存按钮。"),
    ]
    add_table(doc, ["按钮/控件", "功能说明"], rows, widths=[Cm(4.7), Cm(11.5)])

    add_heading(doc, "八、图形区域说明", 1)
    rows = [
        ("实时波形图 通道1/通道2", "显示选定通道的时间波形，纵轴为相位(rad)或应变模式下的对应量。可以拖拽、滚轮缩放。"),
        ("频谱瀑布图 / 频谱图", "默认显示 FFT 随时间变化的瀑布图；点击“切换频谱图”后显示当前频谱曲线，再点“切换瀑布图”返回。"),
        ("振动定位瀑布图(RMS)", "显示沿光纤位置的振动强弱随时间变化，默认色条 0-6。横轴可在“米”和“通道”之间切换。"),
        ("初始化视角", "恢复 RMS 图默认显示范围，适用于缩放拖拽后快速回到全局视图。"),
        ("切换单帧图 / 切换瀑布图", "在 RMS 瀑布图和单帧定位曲线之间切换。单帧图便于看当前时刻沿线强弱。"),
        ("暂停刷新 / 继续刷新", "暂停或恢复 RMS 图刷新。暂停仅冻结显示，不代表底层采集停止。"),
        ("X轴: 米 / X轴: 通道", "切换 RMS 横轴单位。米适合现场定位；通道适合与原始数据行号核对。"),
        ("光强图", "辅助显示区，纵轴为光强/db，横轴为位置(m)。当前主链路以相位、RMS 和 FFT 显示为主，若该图为空白不一定代表采集异常。"),
    ]
    add_table(doc, ["区域/按钮", "说明"], rows, widths=[Cm(4.4), Cm(11.8)])


def add_performance_and_save(doc):
    add_heading(doc, "九、处理耗时颜色：不能让它变红", 1)
    add_paragraph(
        doc,
        "界面中的“原始数据、数据重构、解缠绕、数据滤波、RMS计算、fft计算、存储耗时”显示每个环节的处理时间。代码按一帧数据对应的时间窗口判断颜色：处理时间小于阈值为蓝色，超过阈值会变红。",
    )
    rows = [
        ("5km / 20k", "约 250 ms", "红色表示某一环节超过 0.25 秒，实时压力很高。"),
        ("10km / 10k", "约 500 ms", "默认档位下任一环节超过 0.5 秒就需要关注。"),
        ("30km / 3.333k", "约 1500 ms", "频率降低，单帧时间窗口变长。"),
        ("50km / 2k", "约 2500 ms", "频率最低，时间窗口最长。"),
    ]
    add_table(doc, ["光纤长度", "单帧处理阈值", "说明"], rows, widths=[Cm(3.2), Cm(3.0), Cm(10.0)])
    add_note(
        doc,
        "红色意味着什么",
        "红色不是普通提示，而是上位机处理不过来的信号。若持续红色，后续显示会滞后，保存可能跟不上，甚至影响实时判断。现场应优先降低数据量。",
        fill="FCE4D6",
        border="C00000",
    )
    rows = [
        ("原始数据变红", "PCIE 读取或数据搬运压力大。检查硬件连接、降低量程数据压力或确认后台程序占用。"),
        ("数据重构变红", "监测范围太大、抽取太小或差分重构压力大。优先提高抽取系数、缩小监测范围。"),
        ("解缠绕变红", "相位处理压力大。缩小监测范围或提高抽取系数。"),
        ("数据滤波变红", "滤波器处理行数过多。提高抽取系数、缩小范围，或确认是否必须开启滤波。"),
        ("RMS计算变红", "定位图计算或刷新压力大。可暂停刷新、切换显示、提高抽取。"),
        ("fft计算变红", "FFT 输入通道数据过多或刷新压力大。切换显示模式或减少实时负载。"),
        ("存储耗时变红", "磁盘写入慢或保存全部数据量过大。优先停止保存全部、换高速磁盘、提高抽取或缩小范围。"),
    ]
    add_table(doc, ["红色位置", "处理建议"], rows, widths=[Cm(3.6), Cm(12.6)])

    add_heading(doc, "十、保存与回看数据", 1)
    add_paragraph(
        doc,
        "保存功能分为“保存单点”和“保存全部”。保存模块会在用户选择的目录下自动创建包含参数的文件夹，文件夹名记录频率、抽取、差分距离、起止通道、rows、cols、pitch、interval、startM、endM，方便后续回看时解析。",
    )
    rows = [
        ("保存单点", "保存一个通道的时间序列。适合长期跟踪某个位置。", "文件夹以 data_ 开头；回看时选择该文件夹。"),
        ("保存全部", "保存整个监测范围的二维数据。适合事后做空间-时间瀑布图分析。", "文件夹以 data_all_ 开头；回看时选择文件夹和其中一个 bin。"),
        ("结束保存单点", "停止单点保存。", "不要直接关闭软件代替结束保存。"),
        ("结束保存全部", "停止保存全部。", "保存全部数据量大，结束后再拷贝文件更稳妥。"),
        ("磁盘剩余空间", "显示当前磁盘可用空间。", "低于 5GB 变橙色，低于 1GB 变红色；保存模块会预留约 100MB 安全空间。"),
        ("存储状态", "显示保存全部状态。", "未保存为红点；保存全部中为绿色状态。"),
    ]
    add_table(doc, ["功能", "用途", "注意事项"], rows)

    add_heading(doc, "十一、查看单点和查看保存全部", 1)
    rows = [
        ("查看单点", "点击后选择单点保存文件夹。窗口显示单点波形图和单点 FFT 频谱图。"),
        ("查看保存全部", "点击后选择保存全部文件夹，再选择该文件夹内的 .bin 文件。窗口显示保存全部瀑布图、所选通道波形和所选通道 FFT。"),
        ("FFT通道", "在保存全部回看窗口中输入要分析的绝对通道号。软件会自动换算为保存文件中的相对行号。"),
        ("查看通道", "按当前 FFT通道 重新读取该通道波形与 FFT。"),
        ("频段", "保存全部回看窗口中切换 FFT 横轴频段：0.001-1Hz、1-10Hz、10-100Hz。"),
        ("关闭", "关闭回看窗口，不影响主界面采集或保存状态。"),
    ]
    add_table(doc, ["控件", "说明"], rows, widths=[Cm(3.6), Cm(12.6)])


def add_config_and_faults(doc):
    add_heading(doc, "十二、保存配置与读取配置", 1)
    add_paragraph(
        doc,
        "“保存配置”会生成一个 txt 文件，记录当前关键参数；“读取配置”会把这些参数重新写回软件。配置项包括 pulse_frequency、extract_count、start_save、end_save、monitorPosition、monitorPosition1、monitorPosition2、singleSaveRow、differential_distance、enable_diff_phase、enable_filter、high_pass_cutoff_hz、numThreads。",
    )
    rows = [
        ("换班或交接", "保存一份当前参数，下一班直接读取，减少手动输入错误。"),
        ("长期监测", "把稳定配置固化下来，异常后可快速恢复。"),
        ("调试对比", "不同工况保存不同配置文件，便于回退。"),
        ("注意", "读取配置后仍要人工检查差分距离对应的物理标距是否不小于 3.2m。"),
    ]
    add_table(doc, ["场景", "建议"], rows, widths=[Cm(3.4), Cm(12.8)])

    add_heading(doc, "十三、常见问题处理", 1)
    rows = [
        ("通道距离显示不符合预期", "先确认光纤长度和抽取系数。通道间距 = 当前点距 × 抽取系数。"),
        ("差分距离设为 8 后长距离下空间分辨率变大", "这是正常现象。30km 每点 1.2m，8 点是 9.6m；50km 每点 2m，8 点是 16m。长距离要人工改小，但不能低于 3.2m。"),
        ("处理耗时持续红色", "提高抽取系数、缩小监测范围、暂停不必要显示、停止保存全部或换更快磁盘。"),
        ("保存全部文件很大", "保存全部是二维全范围数据，数据量远大于单点保存。只需要某个位置时请用保存单点。"),
        ("查看保存全部时 FFT 通道不对", "确认输入的是保存时的绝对通道号；当前版本已按 startChannel 自动转换为文件内部行号。"),
        ("图形被缩放后看不到数据", "点击对应图上的“初始化”或“初始化视角”。"),
        ("光强图为空", "当前主要数据链路是相位/RMS/FFT；光强图为空不等同于 PCIE 无数据，应结合实时波形和处理耗时判断。"),
    ]
    add_table(doc, ["问题", "处理方法"], rows, widths=[Cm(4.8), Cm(11.4)])

    add_heading(doc, "十四、现场参数检查清单", 1)
    rows = [
        ("□", "光纤长度已确认", "与现场实际量程一致。"),
        ("□", "抽取系数已确认", "数据量能承受，处理耗时不红。"),
        ("□", "起止通道已确认", "范围覆盖目标区域，不过度扩大。"),
        ("□", "差分距离已确认", "物理标距不小于 3.2m。"),
        ("□", "是否滤波/高通截止已确认", "默认勾选滤波，截止 1.00Hz；特殊工况另行确认。"),
        ("□", "实时波形/RMS/FFT显示正常", "图形有数据，坐标范围正确。"),
        ("□", "处理耗时均未持续红色", "若红色，先降低数据量再开始正式监测。"),
        ("□", "保存状态和磁盘空间正常", "保存前确认磁盘空间充足。"),
        ("□", "配置已保存", "正式运行前建议保存当前配置。"),
    ]
    add_table(doc, ["检查", "项目", "判定标准"], rows, widths=[Cm(1.4), Cm(4.8), Cm(10.0)])


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_quick_start(doc)
    add_parameter_sections(doc)
    add_controls(doc)
    add_performance_and_save(doc)
    add_config_and_faults(doc)

    core_props = doc.core_properties
    core_props.title = "Real_Das 软件使用说明书"
    core_props.subject = "分布式光纤振动监测软件现场用户手册"
    core_props.author = "Codex"
    core_props.keywords = "Real_Das, DAS, 光纤长度, 抽取系数, 差分距离, 空间分辨率"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
