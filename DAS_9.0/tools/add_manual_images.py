from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "Real_Das_用户使用说明书.docx"
OUT_DOCX = ROOT / "docs" / "Real_Das_用户使用说明书_带图片.docx"
ASSET_DIR = ROOT / "docs" / "manual_assets"
SCREENSHOT = ASSET_DIR / "user_full_ui.png"
BACKUP = ROOT / "docs" / "Real_Das_用户使用说明书_加图前备份.docx"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for candidate in candidates:
        path = Path(candidate)
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


FONT = font(28)
FONT_SMALL = font(22)
FONT_BOLD = font(30, True)


def crop_image(src: Image.Image, box: tuple[int, int, int, int], name: str) -> Image.Image:
    img = src.crop(box).convert("RGB")
    path = ASSET_DIR / name
    img.save(path, quality=95)
    return img


def annotate(
    img: Image.Image,
    marks: list[tuple[int, int, int, int, str]],
    name: str,
    title: str | None = None,
) -> Path:
    img = img.convert("RGB")
    pad_top = 76 if title else 18
    canvas = Image.new("RGB", (img.width + 36, img.height + pad_top + 18), "white")
    draw = ImageDraw.Draw(canvas)
    if title:
        draw.text((18, 18), title, fill=(28, 62, 98), font=FONT_BOLD)
    canvas.paste(img, (18, pad_top))
    draw = ImageDraw.Draw(canvas)
    for idx, (x1, y1, x2, y2, label) in enumerate(marks, 1):
        x1 += 18
        x2 += 18
        y1 += pad_top
        y2 += pad_top
        draw.rounded_rectangle((x1, y1, x2, y2), radius=8, outline=(220, 40, 40), width=5)
        badge_r = 18
        bx, by = x1 + 8, max(y1 - 28, pad_top + 4)
        draw.ellipse((bx, by, bx + badge_r * 2, by + badge_r * 2), fill=(220, 40, 40))
        draw.text((bx + 10, by + 2), str(idx), fill="white", font=FONT_SMALL)
        tx = min(x2 + 12, canvas.width - 360)
        ty = max(y1, pad_top + 6)
        draw.rounded_rectangle((tx, ty, tx + 330, ty + 42), radius=8, fill=(255, 245, 230), outline=(220, 160, 90), width=2)
        draw.text((tx + 12, ty + 7), label, fill=(45, 65, 90), font=FONT_SMALL)
    if canvas.width < 1280:
        wide = Image.new("RGB", (1280, canvas.height), "white")
        wide.paste(canvas, ((1280 - canvas.width) // 2, 0))
        canvas = wide
    out = ASSET_DIR / name
    canvas.save(out, quality=95)
    return out


def make_flowchart() -> Path:
    w, h = 1760, 430
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 30), "软件内部数据流程", fill=(28, 62, 98), font=FONT_BOLD)
    steps = [
        ("PCIE采集", "FPGA上传原始点"),
        ("数据重构", "rebuild_data"),
        ("相位解缠绕", "unwrap"),
        ("数据滤波", "高通滤波"),
        ("RMS/FFT", "定位瀑布图与频谱"),
        ("数据保存", "单点/全部数据"),
    ]
    x, y, bw, bh, gap = 48, 145, 230, 118, 48
    colors = [(232, 243, 255), (238, 248, 239), (255, 249, 230), (245, 240, 255), (238, 249, 255), (255, 241, 241)]
    for i, ((title, sub), color) in enumerate(zip(steps, colors)):
        x0 = x + i * (bw + gap)
        draw.rounded_rectangle((x0, y, x0 + bw, y + bh), radius=20, fill=color, outline=(110, 150, 190), width=3)
        draw.text((x0 + 34, y + 24), title, fill=(18, 51, 84), font=FONT_BOLD)
        draw.text((x0 + 34, y + 72), sub, fill=(65, 83, 102), font=FONT_SMALL)
        if i < len(steps) - 1:
            ax1, ay = x0 + bw + 8, y + bh // 2
            ax2 = x0 + bw + gap - 8
            draw.line((ax1, ay, ax2, ay), fill=(88, 120, 160), width=5)
            draw.polygon([(ax2, ay), (ax2 - 18, ay - 12), (ax2 - 18, ay + 12)], fill=(88, 120, 160))
    draw.rounded_rectangle((46, 318, 1715, 380), radius=18, fill=(248, 251, 255), outline=(160, 190, 220), width=2)
    draw.text(
        (72, 334),
        "处理耗时若连续变红，说明当前数据量或计算量超过上位机承受能力，应提高抽取系数、缩小监测范围或减少不必要显示。",
        fill=(45, 65, 90),
        font=FONT_SMALL,
    )
    out = ASSET_DIR / "fig_02_data_flow.png"
    img.save(out, quality=95)
    return out


def make_assets() -> dict[str, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    src = Image.open(SCREENSHOT).convert("RGB")
    assets: dict[str, Path] = {}

    full = src.copy()
    full.thumbnail((1700, 900))
    assets["一、比较细致的操作顺序"] = annotate(
        full,
        [
            (1130, 330, 1290, 385, "光纤长度"),
            (1130, 380, 1290, 430, "抽取系数"),
            (1325, 330, 1590, 430, "起止通道"),
            (1260, 430, 1510, 485, "差分距离"),
            (1085, 480, 1285, 550, "保存全部"),
            (55, 65, 195, 105, "回看入口"),
            (0, 540, 1120, 855, "RMS定位图"),
        ],
        "fig_01_operation_order.png",
        "推荐操作顺序总览",
    )

    assets["二、软件内部数据流程"] = make_flowchart()

    panel = crop_image(src, (1700, 475, 2550, 905), "crop_right_panel.png")
    assets["三、光纤长度：先确定量程和基础点距"] = annotate(
        panel,
        [(92, 25, 230, 70, "光纤长度选择"), (228, 25, 292, 70, "确定")],
        "fig_03_fiber_length.png",
        "光纤长度设置位置",
    )

    assets["四、抽取系数：为什么必须抽取"] = annotate(
        panel,
        [(92, 70, 230, 116, "抽取系数"), (228, 70, 292, 116, "确定")],
        "fig_04_extract.png",
        "抽取系数设置位置",
    )

    assets["五、监测起始通道和结束通道"] = annotate(
        panel,
        [(392, 25, 570, 70, "监测起始通道"), (392, 70, 570, 116, "监测结束通道"), (570, 25, 665, 116, "通道间距提示")],
        "fig_05_monitor_channels.png",
        "监测通道范围设置",
    )

    assets["六、差分距离（标距）：空间分辨率设置"] = annotate(
        panel,
        [(300, 116, 505, 160, "差分距离(标距)"), (505, 116, 575, 160, "确定"), (585, 118, 685, 160, "滤波开关")],
        "fig_06_diff_distance.png",
        "差分距离与空间分辨率",
    )

    ch1 = crop_image(src, (980, 175, 1698, 230), "crop_channel_buttons.png")
    assets["七、主界面按钮和控件说明"] = annotate(
        ch1,
        [(10, 5, 250, 50, "当前通道与位置"), (255, 5, 390, 50, "通道输入"), (392, 5, 455, 50, "确定"), (462, 5, 575, 50, "保存单点"), (582, 5, 700, 50, "结束/初始化")],
        "fig_07_main_controls.png",
        "通道面板按钮",
    )

    graph = crop_image(src, (0, 175, 2550, 1320), "crop_graph_areas.png")
    graph.thumbnail((1700, 900))
    assets["八、图形区域说明"] = annotate(
        graph,
        [(35, 35, 1120, 200, "实时波形图"), (1180, 0, 1690, 245, "频谱瀑布图"), (45, 520, 1120, 870, "RMS定位瀑布图"), (1180, 690, 1690, 880, "光强图")],
        "fig_08_graph_areas.png",
        "图形区域对应关系",
    )

    assets["九、处理耗时颜色：不能让它变红"] = annotate(
        panel,
        [(260, 155, 610, 245, "各处理环节耗时"), (595, 245, 720, 285, "保存状态")],
        "fig_09_processing_time.png",
        "处理耗时与保存状态",
    )

    assets["十、保存与回看数据"] = annotate(
        panel,
        [(15, 155, 225, 200, "保存全部/结束保存全部"), (15, 245, 230, 295, "保存/读取配置"), (260, 155, 610, 245, "保存时关注耗时")],
        "fig_10_save_data.png",
        "保存相关按钮",
    )

    top = crop_image(src, (0, 70, 330, 135), "crop_top_buttons.png")
    assets["十一、查看单点和查看保存全部"] = annotate(
        top,
        [(10, 18, 135, 55, "查看单点"), (145, 18, 290, 55, "查看保存全部")],
        "fig_11_replay_buttons.png",
        "数据回看入口",
    )

    assets["十二、保存配置与读取配置"] = annotate(
        panel,
        [(15, 245, 115, 295, "保存配置"), (125, 245, 230, 295, "读取配置"), (92, 25, 230, 116, "配置会记录关键参数")],
        "fig_12_config.png",
        "配置保存与恢复",
    )

    assets["十三、常见问题处理"] = annotate(
        panel,
        [(260, 155, 610, 245, "红色耗时需降负载"), (595, 245, 720, 285, "未保存状态"), (300, 116, 505, 160, "检查标距")],
        "fig_13_troubleshooting.png",
        "常见问题定位位置",
    )

    assets["十四、现场参数检查清单"] = annotate(
        full,
        [
            (1130, 330, 1290, 385, "量程"),
            (1130, 380, 1290, 430, "抽取"),
            (1325, 330, 1590, 430, "范围"),
            (1260, 430, 1510, 485, "标距"),
            (1260, 485, 1510, 550, "滤波/耗时"),
            (1085, 480, 1285, 550, "保存"),
        ],
        "fig_14_checklist.png",
        "现场参数检查位置",
    )
    return assets


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_picture_after(paragraph, image_path: Path, caption: str):
    spacer = insert_paragraph_after(paragraph, "")
    spacer.paragraph_format.space_after = Pt(2)
    pic_para = insert_paragraph_after(spacer, "")
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_para.add_run()
    run.add_picture(str(image_path), width=Inches(6.45))
    cap_para = insert_paragraph_after(pic_para, caption)
    cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap_para.runs[0]
    cap_run.font.size = Pt(9)
    cap_run.font.color.rgb = RGBColor(90, 105, 125)
    cap_para.paragraph_format.space_after = Pt(10)
    return cap_para


def main() -> None:
    if not DOCX.exists():
        raise FileNotFoundError(DOCX)
    if not SCREENSHOT.exists():
        raise FileNotFoundError(SCREENSHOT)
    if not BACKUP.exists():
        shutil.copy2(DOCX, BACKUP)

    assets = make_assets()
    doc = Document(str(DOCX))
    inserted = 0
    for para in list(doc.paragraphs):
        title = para.text.strip()
        if title in assets:
            caption = "图示：" + title.split("、", 1)[-1]
            add_picture_after(para, assets[title], caption)
            inserted += 1
    doc.save(str(OUT_DOCX))
    print(f"inserted={inserted}")
    print(OUT_DOCX)
    print(BACKUP)


if __name__ == "__main__":
    main()
