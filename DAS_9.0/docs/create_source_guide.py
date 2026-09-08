from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\daimaxuexi\DAS_9.0")
OUTPUT = Path(r"D:\daimaxuexi\DAS_9.0_源码注释与流程说明.docx")
ASSET_DIR = ROOT / "docs" / "_source_guide_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
TEXT_GRAY = "555555"


def font(size, bold=False, color=(11, 37, 69)):
    for candidate in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf"):
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size, index=0)
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color=(46, 116, 181), width=5):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = end
    x0, y0 = start
    if abs(x1 - x0) >= abs(y1 - y0):
        direction = 1 if x1 > x0 else -1
        tri = [(x1, y1), (x1 - 18 * direction, y1 - 10), (x1 - 18 * direction, y1 + 10)]
    else:
        direction = 1 if y1 > y0 else -1
        tri = [(x1, y1), (x1 - 10, y1 - 18 * direction), (x1 + 10, y1 - 18 * direction)]
    draw.polygon(tri, fill=color)


def node(draw, box, title, subtitle="", fill=(235, 242, 250), border=(46, 116, 181)):
    draw.rounded_rectangle(box, radius=20, fill=fill, outline=border, width=4)
    x1, y1, x2, y2 = box
    title_font = font(28, bold=True)
    sub_font = font(19, color=(70, 85, 105))
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((x1 + x2 - (title_box[2] - title_box[0])) / 2, y1 + 18), title, font=title_font, fill=(11, 37, 69))
    if subtitle:
        lines = subtitle.split("\n")
        y = y1 + 65
        for line in lines:
            b = draw.textbbox((0, 0), line, font=sub_font)
            draw.text(((x1 + x2 - (b[2] - b[0])) / 2, y), line, font=sub_font, fill=(70, 85, 105))
            y += 28


def canvas(title, subtitle, w=2000, h=1120):
    im = Image.new("RGB", (w, h), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, w, 126), fill=(11, 37, 69))
    d.text((56, 26), title, font=font(40, bold=True), fill="white")
    d.text((58, 80), subtitle, font=font(21), fill=(212, 225, 240))
    return im, d


def overall_diagram():
    im, d = canvas("实时数据处理总链路", "蓝色箭头表示主数据流；虚线/分支表示同一帧的并行消费。", 2200, 1180)
    nodes = [
        (60, 210, 310, 340, "PCIe / FPGA", "双缓冲 DMA 帧"),
        (395, 210, 655, 340, "receive_data", "轮询、分块 C2H\n附带采集元数据"),
        (740, 210, 1000, 340, "Rebuild_data", "解码、空间差分\n抽取、ping-pong"),
        (1085, 210, 1345, 340, "Unwrap", "连续相位解缠\n可选时间差分"),
        (1430, 210, 1690, 340, "Filter", "每行高通滤波\n提取双监听行"),
    ]
    for b in nodes:
        node(d, b[:4], b[4], b[5])
    for i in range(len(nodes) - 1):
        draw_arrow(d, (nodes[i][2], 275), (nodes[i + 1][0], 275))
    node(d, (1085, 510, 1345, 640), "Sava_data", "全帧/单行快照\n滚动 .bin 文件", (250, 244, 232), (160, 120, 30))
    draw_arrow(d, (1215, 340), (1215, 510), (160, 120, 30))
    downstream = [
        (250, 820, 540, 965, "rms_calculate", "时间分组 RMS\n定位瀑布/曲线"),
        (770, 820, 1040, 965, "Audio", "实时声卡环形缓冲"),
        (1260, 820, 1530, 965, "FFT_Calculate", "单边幅度频谱"),
        (1750, 820, 2070, 965, "MainWindow", "数据渲染、UI、\n耗时与状态"),
    ]
    for b in downstream:
        node(d, b[:4], b[4], b[5], (241, 248, 243), (46, 130, 100))
    draw_arrow(d, (1560, 340), (395, 820), (46, 130, 100))
    draw_arrow(d, (1560, 340), (905, 820), (46, 130, 100))
    draw_arrow(d, (1560, 340), (1395, 820), (46, 130, 100))
    draw_arrow(d, (540, 892), (1750, 892), (46, 116, 181))
    draw_arrow(d, (1040, 892), (1750, 892), (46, 116, 181))
    draw_arrow(d, (1530, 892), (1750, 892), (46, 116, 181))
    path = ASSET_DIR / "01_overall_pipeline.png"
    im.save(path)
    return path


def threading_diagram():
    im, d = canvas("线程与责任边界", "QThread 将 CPU/IO 重任务从 GUI 线程隔离；信号槽在对象所属线程排队投递。", 2200, 1220)
    rows = [
        (175, "GUI / 主线程", "MainWindow、MainWindowData、MainWindowPlot、SavedDataViewer"),
        (360, "PCIe 线程", "receive_data：1 ms 精确定时器、DMA 读取、硬件确认"),
        (545, "重构线程", "Rebuild_data：列分块并行差分（内部再使用 QtConcurrent）"),
        (730, "相位线程", "Unwrap：解缠、差分相位、保存快照分流"),
        (915, "滤波线程", "Filter：每行 IIR 状态、双监听行、下游分发"),
        (1100, "并行消费线程", "RMS、FFT、Audio、Sava_data：计算、播放、文件 I/O"),
    ]
    y0 = 155
    for idx, (y, label, desc) in enumerate(rows):
        d.rounded_rectangle((70, y, 480, y + 110), 18, fill=(235, 242, 250), outline=(46, 116, 181), width=3)
        d.text((98, y + 17), label, font=font(28, True), fill=(11, 37, 69))
        d.rounded_rectangle((570, y, 2110, y + 110), 18, fill=(249, 251, 253), outline=(183, 201, 221), width=2)
        d.text((605, y + 33), desc, font=font(23), fill=(40, 58, 80))
        if idx < len(rows) - 1:
            draw_arrow(d, (275, y + 110), (275, rows[idx + 1][0]), (100, 130, 165), 4)
    d.text((710, 179), "界面事件、目录对话框、绘图刷新", font=font(18), fill=(86, 100, 115))
    d.text((710, 364), "Raw_data(frame, mode, rows, cols, …)", font=font(18), fill=(86, 100, 115))
    d.text((710, 549), "widget_my_array_Signal(float*, …)", font=font(18), fill=(86, 100, 115))
    d.text((710, 734), "Unwrap_Data_Signal / SaveSnapshot", font=font(18), fill=(86, 100, 115))
    d.text((710, 919), "Filter_my_array / multiAudioRows / FFT", font=font(18), fill=(86, 100, 115))
    path = ASSET_DIR / "02_threading.png"
    im.save(path)
    return path


def algorithm_diagram():
    im, d = canvas("重构、解缠与滤波的逐帧处理", "每个步骤均只处理一帧；配置改变会清空与数据连续性相关的状态。", 2200, 1070)
    nodes = [
        (80, 260, 410, 430, "原始帧", "FPGA 列主序\n4 字节/样本"),
        (505, 260, 850, 430, "buildDiffFromRaw", "取高 16 位有符号相位\n差分距离 D\n每 E 行抽取"),
        (945, 260, 1290, 430, "array_unwrap_matrix_memory", "逐行保持上一帧末值\n将相邻差限制到 [-π, π)"),
        (1385, 260, 1730, 430, "applyDiffPhase（可选）", "首点减前帧末值\n后续点减前一时刻"),
        (1825, 260, 2130, 430, "FilterWorker", "每行独立 Chebyshev I\n八阶高通"),
    ]
    for b in nodes:
        node(d, b[:4], b[4], b[5])
    for i in range(len(nodes) - 1):
        draw_arrow(d, (nodes[i][2], 345), (nodes[i + 1][0], 345))
    d.rounded_rectangle((270, 670, 1930, 880), radius=22, fill=(250, 244, 232), outline=(160, 120, 30), width=3)
    d.text((320, 708), "状态复位条件", font=font(30, True), fill=(100, 73, 0))
    d.text((320, 770), "行数、列数、抽取率、保存范围、差分距离、滤波开关或差分相位开关改变时，"
           "清空对应历史，避免把不连续采集段错误连接。", font=font(24), fill=(70, 58, 30))
    path = ASSET_DIR / "03_algorithm.png"
    im.save(path)
    return path


def save_diagram():
    im, d = canvas("保存与离线回放链路", "实时保存使用快照隔离写入；离线查看使用后台线程读取，期间暂停实时采集。", 2200, 1120)
    nodes = [
        (90, 230, 410, 390, "MainWindow", "选择保存模式\n请求目录"),
        (520, 230, 840, 390, "Sava_data", "建立会话目录\n校验磁盘余量"),
        (950, 230, 1270, 390, "Unwrap 快照", "QByteArray\n整帧或单行"),
        (1380, 230, 1700, 390, ".bin 文件", "滚动写入\n目录名携带元数据"),
        (1785, 230, 2120, 390, "SavedDataViewer", "选择目录/文件\n启动后台读取"),
    ]
    for b in nodes:
        node(d, b[:4], b[4], b[5], (245, 249, 253), (46, 116, 181))
    draw_arrow(d, (410, 310), (520, 310))
    draw_arrow(d, (1270, 310), (1380, 310), (160, 120, 30))
    draw_arrow(d, (1700, 310), (1785, 310), (46, 130, 100))
    draw_arrow(d, (1270, 650), (1110, 390), (160, 120, 30))
    node(d, (950, 650, 1270, 810), "实时数据帧", "解缠后的 float 矩阵", (250, 244, 232), (160, 120, 30))
    node(d, (600, 900, 1600, 1035), "后台离线分析", "解析目录元数据 → 按频带 IIR 带通 → 构建相位瀑布、通道波形和 FFT → GUI 线程渲染", (241, 248, 243), (46, 130, 100))
    draw_arrow(d, (1950, 390), (1100, 900), (46, 130, 100))
    path = ASSET_DIR / "04_save_offline.png"
    im.save(path)
    return path


def render_diagram():
    im, d = canvas("UI 参数与渲染职责划分", "MainWindow 管理交互与参数；MainWindowPlot 创建结构；MainWindowData 填充实时数据。", 2200, 1020)
    node(d, (85, 240, 530, 430), "MainWindow", "按钮、校验、保存对话框\n性能标签、显示模式" )
    node(d, (875, 140, 1335, 330), "ParameterManager", "线程安全键值表\nparameterChanged 广播", (250, 244, 232), (160, 120, 30))
    node(d, (875, 530, 1335, 720), "MainWindowPlot", "初始化曲线/色图\n标题、坐标、色标、交互")
    node(d, (1660, 240, 2110, 430), "MainWindowData", "队列节流、缓存、\nRMS/FFT/音频重绘", (241, 248, 243), (46, 130, 100))
    draw_arrow(d, (530, 335), (875, 235), (160, 120, 30))
    draw_arrow(d, (530, 335), (875, 625))
    draw_arrow(d, (1335, 235), (1660, 335), (160, 120, 30))
    draw_arrow(d, (1335, 625), (1660, 335))
    d.rounded_rectangle((430, 825, 1780, 945), 20, fill=(249, 251, 253), outline=(183, 201, 221), width=2)
    d.text((470, 855), "输出：QCustomPlot 音频曲线、RMS/应变定位图、FFT 曲线或瀑布图；所有绘制回到 GUI 线程。", font=font(23), fill=(40, 58, 80))
    path = ASSET_DIR / "05_ui_rendering.png"
    im.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
                element = margins.find(qn(f"w:{side}"))
                if element is None:
                    element = OxmlElement(f"w:{side}")
                    margins.append(element)
                element.set(qn("w:w"), value)
                element.set(qn("w:type"), "dxa")
    # 标记首行是可重复的表头，便于跨页阅读并改善辅助技术语义。
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = header_pr.find(qn("w:tblHeader"))
    if header_flag is None:
        header_flag = OxmlElement("w:tblHeader")
        header_pr.append(header_flag)
    header_flag.set(qn("w:val"), "true")


def set_image_alt(shape, description):
    """为嵌入式流程图提供 Word 可访问性说明。"""
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description)


def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_spacing(paragraph, before=0, after=6, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_para(doc, text="", style=None, before=0, after=6, line=1.25, bold_label=None):
    p = doc.add_paragraph(style=style)
    set_para_spacing(p, before, after, line)
    if bold_label and text.startswith(bold_label):
        r = p.add_run(bold_label)
        set_run_font(r, 11, True, NAVY)
        r = p.add_run(text[len(bold_label):])
        set_run_font(r, 11, False, "000000")
    else:
        r = p.add_run(text)
        set_run_font(r, 11, False, "000000")
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    set_para_spacing(p, {1: 18, 2: 14, 3: 10}.get(level, 8), {1: 10, 2: 7, 3: 5}.get(level, 4), 1.1)
    r = p.add_run(text)
    set_run_font(r, {1: 16, 2: 13, 3: 12}.get(level, 11), True, BLUE if level <= 2 else DARK_BLUE)
    p.paragraph_format.keep_with_next = True
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p, 2, 8, 1.0)
    r = p.add_run(text)
    set_run_font(r, 9, False, TEXT_GRAY)
    r.italic = True


def add_table(doc, headers, rows, widths=(2500, 6860)):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para_spacing(p, 0, 0, 1.1)
        r = p.add_run(header)
        set_run_font(r, 9.5, True, NAVY)
        set_cell_shading(hdr[i], PALE_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            set_para_spacing(p, 0, 0, 1.15)
            r = p.add_run(value)
            set_run_font(r, 9.3, i == 0, NAVY if i == 0 else "000000")
    set_table_geometry(table, list(widths))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_method_table(doc, items):
    add_table(doc, ["方法 / 函数", "在本文件中的职责与关键行为"], items, (2800, 6560))


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, 9, False, TEXT_GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [(1, 16, BLUE, 18, 10), (2, 13, BLUE, 14, 7), (3, 12, DARK_BLUE, 10, 5)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.text = ""
    r = p.add_run("DAS 9.0 | 源码注释与流程说明")
    set_run_font(r, 9, True, TEXT_GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = ""
    r = p.add_run("内部技术文档  |  第 ")
    set_run_font(r, 9, False, TEXT_GRAY)
    add_page_number(p)
    r = p.add_run(" 页")
    set_run_font(r, 9, False, TEXT_GRAY)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(118)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DAS 9.0")
    set_run_font(r, 14, True, TEXT_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p, 8, 6, 1.0)
    r = p.add_run("源码注释与流程说明")
    set_run_font(r, 30, True, NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p, 2, 24, 1.15)
    r = p.add_run("Qt/C++ 分布式光纤数据采集、处理、可视化与保存模块导读")
    set_run_font(r, 14, False, DARK_BLUE)
    add_table(doc, ["文档属性", "内容"], [
        ("范围", "工程自身的 Qt/C++ 业务源码、接口注释和运行链路；第三方库单独标识为依赖边界。"),
        ("注释方式", "为业务头文件补充 Doxygen 风格接口契约，为对应实现文件添加文件级处理流程说明；不改变可执行逻辑。"),
        ("阅读目标", "理解一帧数据从 PCIe 到图形界面、音频、频谱和二进制文件的生命周期，并能定位每个 .cpp 的职责。"),
        ("生成日期", str(date.today())),
    ])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(p, 20, 0, 1.0)
    r = p.add_run("文档设计：compact_reference_guide + editorial_cover")
    set_run_font(r, 9.5, False, TEXT_GRAY)
    doc.add_page_break()


SECTIONS = [
    {
        "file": "main.cpp",
        "purpose": "应用的组合根。它创建 QApplication、主窗口和十个工作线程，将所有处理模块移动到合适线程，并把实时数据、保存状态、性能统计与关闭顺序连接成完整系统。",
        "flow": "初始化应用和 PCIe；实例化模块；建立原始数据、显示、保存、参数和性能信号；启动线程；在 aboutToQuit 中先阻断 UI 信号、停止采集、退出并等待线程、释放 PCIe 与对象。",
        "methods": [
            ("main", "创建 QApplication 和 MainWindow；初始化 PCIe；构建 receive_data → Rebuild_data → Unwrap → Filter 主链；连接 RMS、FFT、Audio、Sava_data 与 UI；注册 aboutToQuit 的受控清理。"),
            ("aboutToQuit lambda", "通过 cleanupStarted 防重入；调用 beginShutdown，停止 PCIe 定时器，逐个 quit/wait 工作线程，只有确认停止后才 delete 线程并执行 pcie_deinit。"),
        ],
        "note": "线程 5、6 目前被创建并启动但没有业务对象移动进去；这是现有结构中的空闲槽位，文档保留该事实而不修改行为。",
    },
    {
        "file": "receive_data.cpp",
        "purpose": "PCIe 采集端。以精确定时器轮询硬件帧标志，分块读取 C2H DMA 数据，并为后续重构携带当前频率、空间抽取、保存范围和差分距离。",
        "flow": "定时器触发 PCIE_recevie；检查停止/离线查看/缓冲条件；读取寄存器 16；根据 mode 18/28 从对应地址读完整帧；写确认值；发出 Raw_data 与采集耗时。",
        "methods": [
            ("rowsByFrequency", "把频率档位转换为 FPGA 原始帧行数，作为帧大小和元数据的共同来源。"),
            ("readPcieFrame", "将大帧切分为 8 MiB C2H 传输；任一块失败即终止该帧。"),
            ("构造/析构", "读取初始参数，按最大容量分配两块 4 KiB 对齐的 mode 缓冲；析构先停止再释放。"),
            ("stop", "原子置停止标志，并在对象所属线程停止 QTimer，支持跨线程安全调用。"),
            ("onParameterChanged", "频率变更时暂停轮询、写 FPGA 控制寄存器、重算行/帧大小并恢复；其他参数只更新缓存。"),
            ("PCIE_recevie", "热路径：检查 mode，读帧，确认硬件并发出下游信号与耗时。"),
        ],
        "note": "frame_buf_mode18 与 frame_buf_mode28 只在采集线程重用；下游通过 Qt 排队信号接收，必须保证下一帧不会在下游消费前覆盖相关缓冲。",
    },
    {
        "file": "rebuild_data.cpp",
        "purpose": "把 FPGA 的列主序原始字节流转为业务使用的行主序 float 矩阵，并完成空间差分与空间抽取。",
        "flow": "验证边界 → 将 start/end 通道换成原始行 → 选择 mode 专属 ping-pong 输出缓冲 → 并行按列解码/差分 → 发射处理矩阵 → 翻转该 mode 的输出索引。",
        "methods": [
            ("decodeRawPhaseSample", "读取每个四字节样本的高 16 位，按有符号 int16 解释并乘 1/8192 恢复 float 相位。"),
            ("buildDiffFromRaw", "夹紧裁剪范围，计算合法输出行数；按列分任务并行执行 phase[row + D] - phase[row]，每 E 行取一次。"),
            ("Rebuild_data", "当前不拥有 MainWindow，仅保留兼容构造参数。"),
            ("saveRowsToCSV", "无副作用占位接口；实际二进制保存由 Sava_data 负责。"),
            ("Receive_raw_data", "过滤离线/无效输入，选择 mode 18 或 28 的当前 ping-pong 矩阵，执行重构并上报纯算法耗时。"),
        ],
        "note": "QtConcurrent 的 worker 数受 ParameterManager 的 numThreads、CPU 核数和列数共同约束；UI 所见的重构耗时不含下游信号阻塞时间。",
    },
    {
        "file": "unwrap.cpp",
        "purpose": "维持逐空间行的跨帧相位连续性；可选将连续相位变成时间差分，并把保存数据复制为稳定快照。",
        "flow": "检查流配置是否改变 → 必要时清历史 → 原地逐行解缠 → 可选相位差分 → 根据保存状态复制全帧或单行 QByteArray → 分别向滤波/保存支路发信号。",
        "methods": [
            ("Receive_is_filter", "收到滤波开关后，如状态变化便清空解缠状态，避免两个不连续处理段相连。"),
            ("onParameterChanged", "更新监测行；差分相位开关变化时复位差分历史。"),
            ("Receive_rebuild_data", "承接一帧重构数据，检测 rows/cols/extract/start/end/diff 配置变化，执行核心算法并构建保存快照。"),
            ("array_unwrap_matrix_memory", "每行用上一样本比较当前值；当差不在 [-π, π) 时减去对应整数倍 2π；记住本帧末值。"),
            ("applyDiffPhase", "首点减前帧末值，后续点减本帧前一点；首个历史帧的首点置零。"),
            ("resetUnwrapMemory / resetDiffPhaseMemory", "清空连续性状态，作为配置切换和模式切换的安全边界。"),
        ],
        "note": "保存采用复制策略而非传递原始 float 指针，避免重构 ping-pong 缓冲在异步文件线程使用期间被下一帧覆盖。",
    },
    {
        "file": "Filter.cpp",
        "purpose": "对每个空间行执行独立高通 IIR 处理，提取两路监听行，同时把数据分发到 RMS、音频、FFT 和未滤波显示支路。",
        "flow": "检查数据 → 需要时重建每行滤波器 → 按行块并行滤波 → 复制两行监听数据 → 发给 Audio/FFT/UI → 滤波开启则发 RMS 矩阵，否则均匀提取五列用于应变显示。",
        "methods": [
            ("FilterWorker::operator()", "处理一个连续行块；同一行始终使用自己的 IIR 实例，避免滤波历史在空间位置间串扰。"),
            ("createRowFilters", "按 rows/frequency/cutoff 创建八阶 Chebyshev I 高通；夹紧截止频率，规避 IIR 库无效参数断言。"),
            ("clearRowFilters", "删除全部行滤波器并清空银行。"),
            ("Receive_is_filter", "更新滤波开关；重新启用时 reset 每条滤波器历史。"),
            ("onParameterChanged", "处理双监听位置和高通截止频率；截止频率变化使滤波器银行变脏。"),
            ("array_handleSignal", "主热路径：并行滤波、监听行复制、信号分发、非滤波支路列提取与性能计时。"),
        ],
        "note": "由于 IIR 是有状态滤波器，频率、行数、截止频率或开关转换后重置比复用旧状态更可靠。",
    },
    {
        "file": "rms_calculate.cpp",
        "purpose": "把每个空间行的一帧时间采样按固定窗口压缩成 RMS 值，为振动定位显示提供二维时空数据。",
        "flow": "对每个空间行按 zhenshu=1000 个样本分组 → 计算每组 RMS → 按“组号”为外层索引累积所有空间行 → 向 GUI 发出 QList<QVector<float>>。",
        "methods": [
            ("构造/析构", "创建并释放预留定时器；QLineEdit 由界面拥有。"),
            ("setLineEdit", "保存非拥有 UI 指针并显示当前窗口长度。"),
            ("array_handleSignal", "逐行逐窗口计算 RMS，乘以现有显示系数 5，发射数据与耗时。"),
            ("calculateVarianceOptimized", "辅助/预留函数：按组计算方差和。"),
            ("calculateRMS", "累计平方和、求均方并开平方；尾部不足一组时仍安全处理。"),
        ],
        "note": "cols 必须至少覆盖每一个完整的 1000 样本窗口；当前调用链的采样列数为 5000，因此会生成 5 个时间分组。",
    },
    {
        "file": "fft_calculate.cpp",
        "purpose": "将第一路监听行转换为单边幅度谱，供实时 FFT 曲线或频谱瀑布使用。",
        "flow": "限制有效样本数 → 分配 fftwf_complex 输出 → 在全局锁内创建 r2c 计划 → 执行 → 逐点取 sqrt(real²+imag²) → 发信号 → 销毁计划/释放内存。",
        "methods": [
            ("FFT_Calculate", "轻量构造；计划与输出存储按帧创建。"),
            ("computeFFT", "预留复数结果接口，当前仅清空结果，不参与主链。"),
            ("FFT_handleSignal", "主路径：调用 FFTW 单精度 r2c，生成 N/2+1 个幅值并上报耗时。"),
        ],
        "note": "对 FFTW 计划器加锁只覆盖 create/destroy，不覆盖 fftwf_execute，避免不必要地串行化执行阶段。",
    },
    {
        "file": "fftw_guard.cpp",
        "purpose": "为 FFTW 的规划器提供进程级互斥锁。",
        "flow": "函数以静态局部 QMutex 形式延迟创建，所有 FFT 计划创建/销毁调用该同一实例。",
        "methods": [("fftwPlannerMutex", "返回唯一共享互斥锁的引用；C++11 确保局部静态初始化线程安全。")],
        "note": "该文件小但很关键：它把第三方规划器的生命周期同步从 FFT 业务逻辑中抽离。",
    },
    {
        "file": "audio.cpp",
        "purpose": "把监听行的 float 样本以单声道浮点格式写入系统声卡，目标是低延迟而非无损回放。",
        "flow": "收到监听数据 → 按频率确保 QAudioSink 与环形设备存在 → 必要时复位出错 sink → 按 float 字节写入环形缓冲 → QAudioSink 读取并输出。",
        "methods": [
            ("AudioRingBufferDevice::resetBuffer", "按样本边界对齐容量和目标缓存，重置读写位置。"),
            ("enqueue", "写入样本；容量不足时丢弃最旧数据，且将队列压到目标缓存长度。"),
            ("readData", "声卡读取时返回可读样本；不足部分补零，避免噪声或未初始化内存。"),
            ("dropOldestUnlocked / writeUnlocked / readUnlocked", "持锁下维护环形索引，正确处理跨缓冲尾部的两段 memcpy。"),
            ("ensureSinkForFrequency", "采样率改变时销毁旧 sink/设备，以 float 单声道格式新建并设置适度缓存。"),
            ("array_handleSignal", "验证长度、确保 sink、必要时重启故障输出，再将 QVector 原始字节入队。"),
        ],
        "note": "实时音频积压时丢弃旧样本是有意设计：对监测监听而言，当前声学状态比历史回放更有价值。",
    },
    {
        "file": "sava_data.cpp",
        "purpose": "管理单通道和保存全部两种会话：目录选择、元数据命名、磁盘保护、滚动 .bin 文件和 GUI 状态反馈。",
        "flow": "GUI 请求开始 → 工作线程请求目录 → GUI 回传目录 → 首个快照到达时创建含元数据的会话目录 → 检查空间并写 .bin → 达到阈值滚动文件 → 停止时刷新关闭。",
        "methods": [
            ("rowsByFrequencyForSave / meterPerRawPointForFrequency", "保存元数据使用的频率到行数、空间点距映射。"),
            ("checkDiskSpace", "每秒缓存一次卷可用空间；保留 100 MiB 安全余量并把可用 GB 回报 UI。"),
            ("onButtonClicked_saveall / saveonlyone", "互斥地请求开始两类会话；单点模式锁定触发时的监测行。"),
            ("onFolderSelected", "处理取消、设置活动参数、记录根目录并等待首帧的真实 rows/cols。"),
            ("createNewBinFile", "关闭旧流、检查单次写入空间、创建带序号 .bin，设置 Qt 6.5 单精度数据流。"),
            ("initializeSaveSessionLocked", "构建带频率、抽取、差分、通道/距离范围的目录名，创建目录并打开首个文件。"),
            ("Preserve_rebuild_data / snapshot", "旧接口先复制全帧；快照接口根据会话写整帧或锁定单行，并控制文件滚动与 flush。"),
            ("closeFile / end_save / end_save_all", "关闭流、清理活动标志、通知界面状态。"),
        ],
        "note": "目录名承担离线回放元数据载体的角色，SavedDataViewer 会解析其中的 freq、rows、cols、pitch、interval、起止距离等字段。",
    },
    {
        "file": "saveddataviewer.cpp",
        "purpose": "离线查看保存目录：解析元数据、在后台读取二进制数据、计算相位瀑布/通道波形/FFT，并在 GUI 线程构建三张图。",
        "flow": "选择文件夹或 .bin → 解析目录和实际文件尺寸 → 设置 saved_data_viewer_busy → QtConcurrent 后台读取/带通/抽样/FFT → QFutureWatcher 回 GUI 线程更新图表和说明 → 清除 busy。",
        "methods": [
            ("parseSavedFolderMeta / effectiveMetaForBin", "从目录名和文件大小推断频率、行列、空间标定；真实 bin 帧数可修正目录声明的 rows。"),
            ("sortedBinFiles / inferActualRowsForBin", "选择顺序稳定的 .bin，并根据帧字节数验证有效行数。"),
            ("applySaveAllBandpass", "对离线全量数据的每个通道应用用户选择频带的带通处理。"),
            ("readSaveAllBin / readChannelAnalysis", "匿名后台核心：流式读取、时间/距离抽样、生成瀑布单元、指定通道波形及 FFT。"),
            ("构造/析构", "创建三张 QCustomPlot、通道/频带控件和专用线程池。"),
            ("loadSinglePointFolder / loadSaveAllFolder / loadSaveAllBinFile", "三种入口统一为目录验证、元数据装载和后台任务启动。"),
            ("startSaveAllBinLoad", "后台生成相位瀑布；完成回调只在 GUI 线程创建色图、色标和文本说明。"),
            ("startSaveAllChannelLoad", "后台重新读取选中行，刷新波形和 FFT。"),
            ("updateFftBandRange / onFftBandChanged", "按频带控制 X/Y 轴；频带变更时可重新运行带通瀑布任务。"),
        ],
        "note": "该文件包含较多匿名命名空间的算法辅助函数；它们不对外暴露，但决定离线视图的数据裁剪、滤波、颜色范围和错误处理质量。",
    },
    {
        "file": "mainwindowutils.cpp",
        "purpose": "提供主窗口和图表模块共用的小型确定性函数：性能颜色、距离换算和图表视觉样式。",
        "flow": "上层将频率、抽取率、行数和耗时传入；本文件返回统一的物理换算结果或直接对 QLabel/QCustomPlot 应用样式。",
        "methods": [
            ("updateTimeLabel", "以 5000/frequency 的毫秒预算判断红/蓝样式，并避免文本或样式未变时重复赋值。"),
            ("getEndIntBySelectedText", "识别长度选择控件中的公里数字。"),
            ("meterPerRawPoint", "2000 Hz 为 2.0 m，3333 Hz 为 1.2 m，其余默认 0.4 m。"),
            ("meterPerExtractedChannel / channelToMeter", "把原始点距乘以抽取率，并把零基通道转换成物理距离。"),
            ("calculateMaxX", "计算图表距离横轴右边界。"),
            ("setPlotBackground", "设置白底、轴/网格颜色、子网格、边距和图例样式。"),
        ],
        "note": "空间标定集中在这里和保存模块的同类辅助函数中；若物理标定规则调整，应同步审阅实时显示与离线元数据。",
    },
    {
        "file": "mainwindowplot.cpp",
        "purpose": "创建和重置 QCustomPlot 的结构：标题、轴、色图、色标、曲线、渐变和拖拽缩放交互。",
        "flow": "各 init 函数先清理旧 plottable/布局装饰，再设置标题和轴，创建曲线或 QCPColorMap/ColorScale，最后打开必要交互并触发 queued replot。",
        "methods": [
            ("setPlotTitle / clearPlotLayoutDecorations", "清除旧标题或色标，防止模式切换后布局残留和重复元素。"),
            ("getDefaultGradient / getStrainGradient / getIntensityGradient", "为 RMS/FFT、正负应变和光强分别提供稳定颜色语义。"),
            ("initAudioPlot", "建立双 graph 结构，实际使用第二条曲线；设置时间 ticker、拖拽和快速折线绘制。"),
            ("initRMSWaterPlot", "创建距离-时间 RMS 色图及底部色标，初始范围 0~6。"),
            ("initRMSLinePlot", "创建 RMS 或应变定位曲线，可选通道或距离横轴。"),
            ("RMStoStrainWaterPlot / StrainToRMSWaterPlot", "复用已有色图，切换标题、单位、色标、数据范围和渐变。"),
            ("initFFTWaterPlot / initFFTSpectrumPlot", "分别创建频率-时间瀑布和单帧幅度曲线。"),
            ("initIntensityPlot", "创建光强色图和灰阶色标。"),
        ],
        "note": "此处不接收信号数据，因而与 mainwindowdata.cpp 的职责刻意分离：前者定义画布，后者填充帧。",
    },
    {
        "file": "mainwindowdata.cpp",
        "purpose": "实时显示的数据消费者。它对音频采用队列节流，对 RMS/应变保留最新帧，对 FFT 支持曲线或瀑布，并确保绘图在 UI 线程发生。",
        "flow": "信号到达 → 依据类型缓存或按帧更新 → 根据显示模式选曲线/色图 → 受 QElapsedTimer 限流 replot；模式切换时从缓存重建可视状态。",
        "methods": [
            ("setAudioPlot(s) / setRMSWaterPlot / setFFTWaterPlot", "绑定绘图对象并初始化相应缓存/队列。"),
            ("shutdown", "停止音频渲染定时器、清队列、断开图表引用，防关闭后访问。"),
            ("resetAudioTimeline / clearAudioChannelData", "清除时间键和曲线，恢复指定音频通道的初始状态。"),
            ("shouldReplot", "以每类图独立的最小间隔限制 queued replot。"),
            ("buildRmsXAxis / getPhaseToMicroStrainScale", "构造通道或距离横轴，并完成相位到微应变的物理比例换算。"),
            ("renderAudioToPlot", "按像素预算下采样显示点，维护连续时间轴和历史上限，按可见数据自动缩放 Y。"),
            ("onReceiveMultiAudioData / onAudioRenderTick", "先将行数据入队，定时按积压量预算取出并批量重绘，避免高频信号直接重绘。"),
            ("onReceiveRMSData / onReceiveStrainData", "缓存最后一帧；曲线模式调用 renderRmsLineFrame，瀑布模式批量写色图并按模式调整色标。"),
            ("onReceiveFFTData", "曲线模式用 N/2 频率范围建 keys；瀑布模式向频率-时间色图追加一行。"),
        ],
        "note": "音频显示具有“跟随尾部”逻辑：用户拖回历史区域时，系统不会强行把 X 轴拉回最新时间。",
    },
    {
        "file": "mainwindow.cpp",
        "purpose": "应用的 GUI 协调层：读取 .ui、创建动态面板、校验用户输入、更新 ParameterManager、控制保存和离线查看、选择渲染模式，并显示各阶段耗时。",
        "flow": "构造阶段：主题→模块→控件指针→信号槽→图表工具；运行阶段：用户事件写参数或发保存信号，工作线程结果回到 onReceive* 槽，后者更新数据模块与 UI。",
        "methods": [
            ("beginShutdown / 析构", "进入关闭保护，停止显示模块，防止后台排队信号再操作已经销毁的窗口。"),
            ("initModules / initUIWidgetPointers / connectSignalsSlots", "创建绘图/数据模块、收集 .ui 控件、建立 UI/参数/渲染器连接。"),
            ("onReceive*Time", "把各线程耗时发送给 MainWindowUtils::updateTimeLabel，形成处理链性能面板。"),
            ("onReceiveRMSData / StrainData / FFTData / MultiAudioData", "关闭或离线查看时忽略；否则保存元数据并交给 MainWindowData 渲染。"),
            ("参数按钮槽", "onChannel、onExtract、onPulseFrequency、onDiffDistance、onHighPassConfirm 等校验输入，再写 ParameterManager。"),
            ("保存槽", "onSaveOnly、onSaveAll、start/end 通道事件维护按钮状态并发出 Sava_data 所需信号。"),
            ("配置槽", "saveConfigToTextFile/loadConfigFromTextFile 使用文本键值保存并同步 UI。"),
            ("图形工具槽", "syncFftPlotMode、syncRmsPlotMode、RMS 复位/横轴/曲线瀑布/暂停事件协调 Plot 与 Data 两层。"),
            ("openSavedDataViewer", "打开非模态离线浏览对话框，按单点或全量模式选择目录/文件。"),
            ("主题与布局辅助", "applyUiTheme、setupAudioPanels、updateOverlayButtonsGeometry、DPI 缩放和通道距离标签。"),
        ],
        "note": "MainWindow 不是算法计算器；它的核心价值是把用户意图翻译为线程安全参数和信号，并保持 GUI 仅在主线程被更新。",
    },
]


def add_section(doc, index, section):
    add_heading(doc, f"{index:02d}. {section['file']}", 1)
    add_para(doc, section["purpose"], before=0, after=7)
    add_heading(doc, "处理流程", 2)
    add_para(doc, section["flow"], before=0, after=8)
    add_heading(doc, "方法与作用", 2)
    add_method_table(doc, section["methods"])
    add_heading(doc, "实现提示", 2)
    add_para(doc, section["note"], before=0, after=6, bold_label="注意：" if section["note"].startswith("注意：") else None)


def build_document():
    figures = [overall_diagram(), threading_diagram(), algorithm_diagram(), save_diagram(), render_diagram()]
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    add_heading(doc, "阅读范围与注释策略", 1)
    add_para(doc, "本次注释覆盖 CMakeLists.txt 中由 Real_Das 目标直接编译的业务 Qt/C++ 源文件及其接口头文件。改动只增加注释，不改变函数签名、数据布局、线程归属、算法公式、Qt 连接方式或构建选项。", after=8)
    add_table(doc, ["类别", "处理范围与理由"], [
        ("业务实现", "main、采集、重构、相位、滤波、RMS、FFT、音频、保存、离线查看、主窗口和图表相关 .cpp 已补充文件级流程说明。"),
        ("业务接口", "对应 .h 的公共槽、关键私有辅助函数和 C/CUDA 边界接口已补充 Doxygen 风格职责、参数和线程/所有权说明。"),
        ("第三方库", "qcustomplot.cpp/.h、iir1-master 与 FFTW 头文件保持原样，避免在可升级的上游代码中混入项目注释；文档说明其在系统中的调用边界。"),
        ("自动生成文件", "build/ 下的 moc、qrc、编译器识别文件不改动，因为下次构建会重新生成。"),
    ])

    add_heading(doc, "一帧数据的生命周期", 1)
    picture = doc.add_picture(str(figures[0]), width=Inches(6.35))
    set_image_alt(picture, "实时主数据链路：PCIe 采集、重构、解缠、滤波，以及 RMS、音频、FFT、保存和界面分支。")
    add_caption(doc, "图 1  实时主链及四个并行消费支路")
    add_para(doc, "一帧原始数据在 PCIe 线程读入后，先被重构为可处理的 float 差分矩阵。解缠模块将其变为跨帧连续相位，同时为保存模块制作独立快照。随后滤波模块根据用户开关把同一帧导向 RMS 定位、双路音频、单路 FFT，以及未滤波应变显示。主窗口只消费结果，不直接参与重算法。", after=9)

    add_heading(doc, "线程模型与排队信号", 1)
    picture = doc.add_picture(str(figures[1]), width=Inches(6.35))
    set_image_alt(picture, "线程责任图：GUI、PCIe、重构、相位、滤波和并行消费线程的对象与信号流。")
    add_caption(doc, "图 2  业务对象的主要线程归属与数据交接")
    add_para(doc, "信号跨线程时由 Qt 的队列连接投递到接收者所属线程。关闭流程的顺序因此重要：先让 MainWindow 拒绝新绘制，再停止 receive_data 的定时器，最后退出并等待工作线程，避免对象被销毁后仍有排队回调。", after=9)

    add_heading(doc, "核心算法段的连续性约束", 1)
    picture = doc.add_picture(str(figures[2]), width=Inches(6.35))
    set_image_alt(picture, "算法流程图：原始帧经过空间差分、相位解缠、可选时间差分和高通滤波。")
    add_caption(doc, "图 3  原始字节、空间差分、解缠、时间差分和高通滤波")
    add_para(doc, "空间差分与空间抽取决定输出矩阵的行数；解缠和差分相位则依赖上一样本或上一帧末值。任何会改变数据语义的配置变化都必须重置相应历史。本工程已在这些边界调用 resetUnwrapMemory、resetDiffPhaseMemory 或重建滤波器组。", after=9)

    add_heading(doc, "保存、离线回放与图形渲染", 1)
    picture = doc.add_picture(str(figures[3]), width=Inches(6.35))
    set_image_alt(picture, "保存和离线回放流程图：目录选择、快照、滚动 bin 文件、后台离线分析和可视化。")
    add_caption(doc, "图 4  保存会话与离线回放")
    picture = doc.add_picture(str(figures[4]), width=Inches(6.35))
    set_image_alt(picture, "UI 渲染职责图：MainWindow、ParameterManager、MainWindowPlot 和 MainWindowData 之间的边界。")
    add_caption(doc, "图 5  GUI 参数、绘图结构和实时渲染的职责分离")
    add_para(doc, "离线查看置位 saved_data_viewer_busy 后，实时采集链会短路返回；这避免大文件解析、带通和 FFT 与 DMA 采集争用资源。渲染层采用 MainWindowPlot/ MainWindowData 分离：前者只初始化画布，后者只消费实时数据并节流 replot。", after=9)
    doc.add_page_break()

    add_heading(doc, "逐文件流程与方法说明", 1)
    add_para(doc, "下列章节按运行链路排列。每个条目都列出文件在系统中的位置、主要处理步骤、关键方法及维护时需要保留的约束。", after=8)
    for index, section in enumerate(SECTIONS, start=1):
        add_section(doc, index, section)
        if index != len(SECTIONS):
            doc.add_page_break()

    add_heading(doc, "第三方与非 .cpp 依赖边界", 1)
    add_table(doc, ["文件/目录", "角色与维护策略"], [
        ("qcustomplot.cpp / qcustomplot.h", "第三方绘图库。业务模块通过 QCustomPlot、QCPColorMap、QCPColorScale 和 QCPGraph 建图；保持原文件不变，避免升级与许可追踪困难。"),
        ("iir1-master", "第三方 IIR 滤波库。Filter 使用 Iir::ChebyshevI::HighPass<8>；SavedDataViewer 的离线带通也基于其滤波能力。"),
        ("fftw3.h / fftw-3.3.5-dll64", "FFTW API、导入库和 DLL。fft_calculate.cpp 只使用单精度 fftwf_ 系列并通过 fftw_guard.cpp 同步规划器生命周期。"),
        ("pcie_fun.c / pcie_fun.h", "硬件 C 接口。头文件已补充方向、控制寄存器和初始化/释放含义；具体设备协议仍应以 FPGA/驱动文档为准。"),
        ("my.cu / my.h", "CUDA 向量加法边界接口；当前主实时链未直接调用。"),
        ("build/", "CMake/Qt 自动生成与构建输出；不应手工注释或修改。"),
    ])

    add_heading(doc, "维护检查清单", 1)
    add_table(doc, ["变更类型", "建议同时核对的代码位置"], [
        ("新增频率档位", "receive_data::rowsByFrequency、sava_data::rowsByFrequencyForSave、MainWindow::getRowsByFrequency 与空间点距映射。"),
        ("改变空间标定", "MainWindowUtils::meterPerRawPoint、保存目录的 pitch/interval 生成及 SavedDataViewer 的元数据解析。"),
        ("改变帧布局", "readPcieFrame、decodeRawPhaseSample、buildDiffFromRaw、保存帧字节数计算、离线 readSaveAllBin。"),
        ("改变解缠/滤波参数", "对应状态复位条件；不能把旧配置下的 IIR 或相位历史带到新配置帧。"),
        ("新增图形模式", "MainWindow 的 UI 状态、MainWindowPlot 的结构初始化、MainWindowData 的数据填充与重绘节流应成套更新。"),
        ("改变保存格式", "Sava_data 的会话元数据、SavedDataViewer 的解析/实际帧数推断和离线兼容逻辑应同步演进。"),
    ])

    doc.core_properties.title = "DAS 9.0 源码注释与流程说明"
    doc.core_properties.subject = "Qt/C++ 实时数据采集、处理、可视化和保存模块导读"
    doc.core_properties.author = "Codex"
    doc.core_properties.comments = "Generated from the current DAS_9.0 source tree; source changes are comment-only."
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
