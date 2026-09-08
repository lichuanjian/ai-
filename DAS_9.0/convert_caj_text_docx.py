from pathlib import Path
import sys
import os
import re
import pymupdf
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

repo = Path(__file__).parent / "caj2pdf"
sys.path.insert(0, str(repo))
os.chdir(repo)
from cajparser import CAJParser

src_dir = Path("E:/")
dst_dir = Path("D:/daimaxuexi")
work_dir = dst_dir / "_caj_work"
work_dir.mkdir(exist_ok=True)

def make_docx(pdf_path: Path, out_path: Path, title: str):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '宋体')
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.name = "黑体"
    run._element.rPr.rFonts.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia', '黑体')
    run.font.size = Pt(16)
    doc.add_paragraph("说明：本文件由 CAJ 转换为可编辑 Word 文档。原文中的复杂公式、图表和版式可能需要人工校对。")
    pdf = pymupdf.open(str(pdf_path))
    for page_no, page in enumerate(pdf, start=1):
        if page_no > 1:
            doc.add_page_break()
        hp = doc.add_paragraph()
        hr = hp.add_run(f"第 {page_no} 页")
        hr.bold = True
        hr.font.size = Pt(9)
        hr.font.color.rgb = __import__('docx').shared.RGBColor(100, 100, 100)
        text = page.get_text("text")
        text = text.replace("\x00", "")
        for line in text.splitlines():
            line = re.sub(r"[ \t]+", " ", line).strip()
            if line:
                para = doc.add_paragraph(line)
                para.paragraph_format.first_line_indent = Inches(0.3)
    pdf.close()
    doc.save(str(out_path))

results = []
for src in sorted(src_dir.glob("*.caj")):
    pdf = work_dir / f"{src.stem}.pdf"
    out = dst_dir / f"{src.stem}.docx"
    try:
        if not pdf.exists() or pdf.stat().st_size < 10000:
            CAJParser(str(src)).convert(str(pdf))
        make_docx(pdf, out, src.stem)
        results.append((src.name, "ok", out.name, out.stat().st_size))
    except Exception as exc:
        results.append((src.name, "error", type(exc).__name__, str(exc)))

for row in results:
    print("\t".join(map(str, row)))
